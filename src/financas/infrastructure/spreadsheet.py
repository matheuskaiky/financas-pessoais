"""Reader of the old workbook (CLAUDE.md 13.1): rows normalised into codes and cents.

Only the source sheets are read (``fato_transacoes``, ``dim_contas``, ``dim_categorias``); the
derived ones are never touched. The Portuguese wording of the old workbook (origin and kind names,
statement payments, Pix, the investment sweep) is a data dialect recognised here, so the planner in
``application/imports`` stays free of it. ``openpyxl`` is an optional dependency.
"""

import datetime as dt
import re
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from financas.application.imports.model import (
    LegacyAccountRow,
    LegacyRow,
    LegacyWorkbook,
    Origin,
    RowKind,
    TransferHint,
)
from financas.domain.errors import DomainError
from financas.domain.money import YearMonth
from financas.domain.services.text import clean_text, normalize_search

FACTS_SHEET = "fato_transacoes"
ACCOUNTS_SHEET = "dim_contas"
CATEGORIES_SHEET = "dim_categorias"

_ORIGINS = {"conta corrente": Origin.CHECKING, "cartao de credito": Origin.CARD}
_KINDS = {
    "despesa": RowKind.EXPENSE,
    "receita": RowKind.INCOME,
    "estorno": RowKind.REFUND,
    "transferencia": RowKind.TRANSFER,
}
_PAYMENT_MARKERS = (
    "pagamento fatura",
    "pagto cartao credito",
    "pagamento recebido",
    "pagamento de fatura",
)
_SWEEP_MARKER = "bb rende facil"
_PIX_PREFIXES = (
    "pix - enviado",
    "pix - recebido",
    "transferencia recebida",
    "transferencia enviada",
)
_TIMESTAMP_NAME = re.compile(r"\d{2}/\d{2} \d{2}:\d{2}\s+(.+)$")
_LONG_NUMBER = re.compile(r"^\d{6,}$")
_MAX_NOISE = Decimal("0.0000001")


def read_legacy_workbook(path: Path) -> LegacyWorkbook:
    """Read the source sheets of the old workbook (``DomainError`` on anything unexpected)."""
    try:
        import openpyxl
    except ImportError:
        raise DomainError("IMPORT_READER_MISSING") from None
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [_legacy_row(number, record) for number, record in _records(workbook[FACTS_SHEET])]
        accounts = [_account_row(record) for _, record in _records(workbook[ACCOUNTS_SHEET])]
        categories = [
            clean_text(str(record["categoria"]))
            for _, record in _records(workbook[CATEGORIES_SHEET])
            if record.get("categoria")
        ]
    finally:
        workbook.close()
    return LegacyWorkbook(rows, accounts, categories)


def _records(sheet: Any) -> list[tuple[int, dict[str, Any]]]:
    """``(excel row number, {header: value})`` for every non-empty row under the header."""
    lines = list(sheet.iter_rows(values_only=True))
    if not lines:
        return []
    headers = [str(h).strip() if h is not None else "" for h in lines[0]]
    result: list[tuple[int, dict[str, Any]]] = []
    for offset, values in enumerate(lines[1:], start=2):
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        result.append((offset, dict(zip(headers, values, strict=False))))
    return result


def _legacy_row(number: int, record: dict[str, Any]) -> LegacyRow:
    kind = _KINDS.get(normalize_search(str(record.get("tipo") or "")))
    if kind is None:
        raise DomainError("IMPORT_UNKNOWN_KIND", row=number)
    origin = _ORIGINS.get(normalize_search(str(record.get("origem") or "")))
    if origin is None:
        raise DomainError("IMPORT_UNKNOWN_ORIGIN", row=number)
    description = clean_text(str(record.get("descricao") or ""))
    hint, counterparty = _transfer_hint(kind, origin, description)
    return LegacyRow(
        row_no=number,
        sheet_id=clean_text(str(record.get("id_transacao") or "")),
        date=_date(record.get("data")),
        statement=_year_month(record.get("fatura_ref")),
        institution=clean_text(str(record.get("instituicao") or "")),
        origin=origin,
        kind=kind,
        category=clean_text(str(record.get("categoria") or "")),
        recurring=normalize_search(str(record.get("recorrente") or "")) == "sim",
        description=description,
        place=clean_text(str(record["localidade"])) if record.get("localidade") else None,
        amount_cents=_cents(record.get("valor"), number),
        installment_number=_int_or_none(record.get("parcela_atual")),
        installment_total=_int_or_none(record.get("parcela_total")),
        transfer_hint=hint,
        counterparty=counterparty,
    )


def _date(value: Any) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return dt.date.fromisoformat(str(value)[:10])


def _year_month(value: Any) -> YearMonth | None:
    text = str(value).strip() if value is not None else ""
    return YearMonth.parse(text) if text else None


def _int_or_none(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    return int(float(value))


def _cents(value: Any, row: int) -> int:
    """Float money to cents: float noise is accepted, a third decimal is not."""
    try:
        exact = Decimal(str(value))
    except InvalidOperation:
        raise DomainError("IMPORT_AMOUNT_NOT_CENTS", row=row) from None
    rounded = exact.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if abs(rounded - exact) >= _MAX_NOISE:
        raise DomainError("IMPORT_AMOUNT_NOT_CENTS", row=row)
    return int(rounded * 100)


def _transfer_hint(kind: RowKind, origin: Origin, description: str) -> tuple[TransferHint, str]:
    """What the description of a transfer says: payment, sweep or a Pix with a counterparty."""
    if kind is not RowKind.TRANSFER:
        return TransferHint.NONE, ""
    key = normalize_search(description)
    if any(marker in key for marker in _PAYMENT_MARKERS):
        return TransferHint.STATEMENT_PAYMENT, ""
    if _SWEEP_MARKER in key:
        return TransferHint.SWEEP, ""
    if origin is not Origin.CHECKING:
        return TransferHint.NONE, ""
    return TransferHint.PIX, _counterparty(description)


def _counterparty(description: str) -> str:
    head, _, tail = description.partition(" | ")
    name = head
    if normalize_search(head).startswith(_PIX_PREFIXES):
        found = _TIMESTAMP_NAME.search(tail.strip())
        name = found.group(1) if found else ""
    tokens = [t for t in normalize_search(name).split() if not _LONG_NUMBER.match(t)]
    return " ".join(tokens)


def _account_row(record: dict[str, Any]) -> LegacyAccountRow:
    return LegacyAccountRow(
        institution=clean_text(str(record.get("instituicao") or "")),
        origin=clean_text(str(record.get("origem") or "")),
        nickname=clean_text(str(record.get("apelido") or "")),
        credit_limit_cents=_optional_cents(record.get("limite_cartao")),
        opening_balance_cents=_optional_cents(record.get("saldo_inicial")),
    )


def _optional_cents(value: Any) -> int | None:
    """Raw ``dim_contas`` money: empty, ``PREENCHER`` or unreadable cells are ``None``."""
    if value is None or isinstance(value, str):
        return None
    try:
        return _cents(value, 0)
    except DomainError:
        return None
