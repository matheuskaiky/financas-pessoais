"""A SYNTHETIC copy of the old workbook's layout, for the import tests (no real data)."""

import datetime as dt
from pathlib import Path
from typing import Any

import openpyxl

FACT_HEADERS = [
    "id_transacao", "data", "ano_mes", "fatura_ref", "instituicao", "origem", "categoria",
    "tipo", "recorrente", "descricao", "localidade", "valor", "valor_abs", "parcela_atual",
    "parcela_total", "parcelas_restantes", "valor_comprometido_futuro",
]  # fmt: skip
ACCOUNT_HEADERS = [
    "instituicao", "origem", "apelido", "tipo_conta", "limite_cartao", "saldo_inicial",
    "dia_fechamento", "dia_vencimento", "observacao",
]  # fmt: skip
CATEGORY_HEADERS = [
    "categoria", "grupo", "natureza", "essencial", "tipo_predominante", "orcamento_mensal",
    "realizado_medio", "media_menos_orcamento",
]  # fmt: skip

D = dt.datetime


def fact(
    sheet_id: str,
    when: dt.datetime,
    institution: str,
    origin: str,
    category: str,
    kind: str,
    description: str,
    value: float,
    *,
    statement: str | None = None,
    recurring: str = "Não",
    place: str | None = None,
    number: int | None = None,
    total: int | None = None,
) -> dict[str, Any]:
    return {
        "id_transacao": sheet_id,
        "data": when,
        "ano_mes": when.strftime("%Y-%m"),
        "fatura_ref": statement,
        "instituicao": institution,
        "origem": origin,
        "categoria": category,
        "tipo": kind,
        "recorrente": recurring,
        "descricao": description,
        "localidade": place,
        "valor": value,
        "valor_abs": abs(value),
        "parcela_atual": number,
        "parcela_total": total,
        "parcelas_restantes": None if number is None or total is None else total - number,
        "valor_comprometido_futuro": None,
    }


def default_rows() -> list[dict[str, Any]]:
    """One row of each shape the reader has to understand (synthetic names only)."""
    return [
        fact(
            "r01",
            D(2026, 1, 3),
            "Banco Alfa",
            "Conta Corrente",
            "Alimentação",
            "Despesa",
            "Padaria São João",
            -59.9000000001,
            recurring="Sim",
            place="TERESINA BR",
        ),
        fact(
            "r02",
            D(2026, 1, 5),
            "Banco Alfa",
            "Cartao de Credito",
            "Educacao",
            "Despesa",
            "Curso Online - Parcela 2/6",
            -96.79,
            statement="2026-02",
            number=2,
            total=6,
        ),
        fact(
            "r03",
            D(2026, 1, 6),
            "Banco Alfa",
            "Conta Corrente",
            "Transferencia",
            "Transferencia",
            'Beltrano Silva | Pix enviado: "Cp :00000000-Beltrano Silva"',
            -150.0,
        ),
        fact(
            "r04",
            D(2026, 1, 7),
            "Banco Alfa",
            "Conta Corrente",
            "Transferencia",
            "Transferencia",
            "Pix - Enviado | 01/02 10:30 FULANO DE TAL",
            -300.0,
        ),
        fact(
            "r05",
            D(2026, 1, 8),
            "Banco Alfa",
            "Conta Corrente",
            "Transferencia",
            "Transferencia",
            "Pagamento Fatura - FULANO DE TAL | Pagamento efetuado",
            -1000.0,
        ),
        fact(
            "r06",
            D(2026, 1, 8),
            "Banco Beta",
            "Cartão de Crédito",
            "Transferencia",
            "Transferencia",
            "Pagamento recebido",
            1000.0,
            statement="2026-01",
        ),
        fact(
            "r07",
            D(2026, 1, 9),
            "Banco Alfa",
            "Conta Corrente",
            "Transferencia",
            "Transferencia",
            "BB Rende Fácil",
            500.0,
        ),
        fact(
            "r08",
            D(2026, 1, 10),
            "Banco Alfa",
            "Conta Corrente",
            "Salario",
            "Receita",
            "Salário de Janeiro",
            4000.0,
            recurring="Sim",
        ),
        fact(
            "r09",
            D(2026, 1, 11),
            "Banco Beta",
            "Cartão de Crédito",
            "Estorno",
            "Estorno",
            "Estorno Loja",
            25.5,
            statement="2026-01",
        ),
        fact(
            "r10",
            D(2026, 1, 12),
            "Banco Alfa",
            "Conta Corrente",
            "Transferencia",
            "Transferência",
            "Transferência recebida | 12/01 09:15 FULANO DE TAL",
            200.0,
        ),
        fact(
            "r11",
            D(2026, 1, 13),
            "Banco Alfa",
            "Conta Corrente",
            "Transferencia",
            "Transferencia",
            "Pix - Enviado",
            -10.0,
        ),
    ]


def build_workbook(
    path: Path,
    rows: list[dict[str, Any]] | None = None,
    *,
    accounts: list[list[Any]] | None = None,
    categories: list[list[Any]] | None = None,
) -> Path:
    """Write the workbook: the three source sheets, an empty row and a derived sheet."""
    workbook = openpyxl.Workbook()
    facts = workbook.active
    assert facts is not None
    facts.title = "fato_transacoes"
    facts.append(FACT_HEADERS)
    for row in default_rows() if rows is None else rows:
        facts.append([row.get(h) for h in FACT_HEADERS])
    facts.append([None] * len(FACT_HEADERS))  # an empty row inside the data
    sheet = workbook.create_sheet("dim_contas")
    sheet.append(ACCOUNT_HEADERS)
    for line in accounts if accounts is not None else [
        ["Banco Alfa", "Conta Corrente", "Banco Alfa c/c", "Conta Corrente", None, 304.21, None,
         None, "estimado"],
        ["Banco Beta", "Cartao de Credito", "Banco Beta cartao", "Cartão de Crédito",
         "PREENCHER", None, "25", None, None],
    ]:  # fmt: skip
        sheet.append(line)
    sheet = workbook.create_sheet("dim_categorias")
    sheet.append(CATEGORY_HEADERS)
    for line in categories if categories is not None else [
        ["Alimentação", "Não essencial", "Despesa", "Não", "Despesa", None, 10.0, None],
        ["Educacao", "Essencial", "Despesa", "Sim", "Despesa", None, 0, None],
        ["Salario", "Renda", "Receita", "-", "Receita", None, 0, None],
        [None, None, None, None, None, None, 0, None],
    ]:  # fmt: skip
        sheet.append(line)
    derived = workbook.create_sheet("fluxo_mensal")  # derived: the reader must never read it
    derived.append(["ano_mes", "receitas"])
    derived.append(["2026-01", "not a number"])
    workbook.save(path)
    return path
